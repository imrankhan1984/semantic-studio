"""
================================================================================
FILE: backend/tests/test_snapshots.py
================================================================================

SUMMARY
    Data snapshots in a project, end to end through the API (csv-data-import
    5.3 to 5.7, Section 9, rows D1 to D12): the wizard's preview, the import
    and its report, the snapshot folder, where its data shows and where it
    stops showing, refresh, change and edit the mapping, remove, the
    2,000-row limit on every path, and the project zip. Stage B (5.8, 5.9,
    rows X1 and X3): a workbook imported by sheet and header row with its
    CSV copy and refreshed the same way, the limit per sheet, and links
    between snapshots resolving in either order, with the ones that match
    no row reported and caught by the Small template's Points to rule.

BASIC IDEA
    Each test makes a Small-template project (Person, Organization, member
    of), adds the attributes it needs by command, and imports a CSV or a
    workbook built in the test. What is asserted is what the spec's rows check: the statements
    written, the report, the label on the data, and save and reopen.

    A snapshot is not a model change: every test that changes one asserts
    the model's revision, dirty flag and undo history did not move.

INPUTS / INPUT SOURCES
    - The FastAPI app through TestClient, with the client header.

EXPECTED OUTPUT
    - Pass/fail; two `perf` budgets (Section 10).
================================================================================
"""

from __future__ import annotations

import datetime
import gc
import io
import json
import shutil
import tempfile
import time
import zipfile
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from rdflib import Graph, Literal, URIRef
from rdflib.namespace import RDF, RDFS, SKOS, XSD

from app import rml
from app.editing import EditingService, editing_service, project_store
from app.main import app
from app.projects import ProjectStore
from app.store import OntologyStore
from tests.budget import limit_ms

client = TestClient(app, base_url="http://localhost", headers={"X-Semantic-Studio": "1"})

EX = "http://example.org/shop#"
PERSON = EX + "Person"


def U(local: str) -> URIRef:
    return URIRef(EX + local)


@pytest.fixture(autouse=True)
def _closed():
    yield
    editing_service.close_all()


def _create(template: str = "small") -> str:
    response = client.post(
        "/api/projects", json={"name": "Shop", "template": template, "baseIri": EX, "prefix": "shop"}
    )
    assert response.status_code == 200, response.text
    project = response.json()["id"]
    assert client.post(f"/api/projects/{project}/open").status_code == 200
    return project


def command(pid: str, name: str, **args) -> dict:
    response = client.post(f"/api/projects/{pid}/documents/model/commands", json={"command": name, "args": args})
    assert response.status_code == 200, response.text
    return response.json()


@pytest.fixture
def pid() -> str:
    """A Small project with a birth date (date) and a nickname (text)."""
    project = _create()
    command(project, "CreateDatatypeProperty", label="birth date", domain=PERSON, datatype=str(XSD.date))
    command(project, "CreateDatatypeProperty", label="nickname", domain=PERSON)
    return project


def model_doc(pid: str):
    return editing_service.document(pid, "model")


def oid(pid: str) -> str:
    return f"{pid}-model"


def csv_file(text: str, name: str = "people.csv", encoding: str = "utf-8"):
    return {"file": (name, text.encode(encoding), "text/csv")}


PEOPLE = (
    "id,name,birth date,nickname,org\n"
    "1,Alice Martin,1990-01-02,Al,acme\n"
    "2,Bob Stone,1985-12-31,,acme\n"
    "3,Cy Young,2000-02-29,Cyc,\n"
)


def choices(**columns) -> dict:
    base = {
        "name": {"as": "name"},
        "birth date": {"as": "attribute", "property": EX + "birthDate"},
        "nickname": {"as": "attribute", "property": EX + "nickname"},
        "org": {"as": "relationship", "property": EX + "memberOf"},
    }
    base.update(columns)
    return {"classIri": PERSON, "idColumn": "id", "columns": base}


def import_csv(pid: str, text: str = PEOPLE, chosen: dict | None = None, *, options: dict | None = None,
               name: str = "people.csv", encoding: str = "utf-8", status: int = 200) -> dict:
    response = client.post(
        f"/api/projects/{pid}/data",
        files=csv_file(text, name, encoding),
        data={"choices": json.dumps(chosen or choices()), "options": json.dumps(options or {})},
    )
    assert response.status_code == status, response.text
    return response.json()


def node(pid: str, iri: str) -> dict:
    response = client.get(f"/api/ontologies/{oid(pid)}/node", params={"iri": iri})
    assert response.status_code == 200, response.text
    return response.json()


def sparql(pid: str, query: str) -> dict:
    response = client.post(f"/api/ontologies/{oid(pid)}/sparql", json={"query": query})
    assert response.status_code == 200, response.text
    return response.json()


def validate(pid: str) -> dict:
    response = client.post(f"/api/projects/{pid}/validate")
    assert response.status_code == 200, response.text
    return response.json()


def folder(pid: str, sid: str) -> Path:
    return project_store.folder(pid) / "data" / sid


def data_graph(pid: str, sid: str) -> Graph:
    return Graph().parse(folder(pid, sid) / "data.ttl", format="nt")


def untouched(pid: str, before: tuple) -> None:
    """A snapshot action is not a model change (5.7)."""
    doc = model_doc(pid)
    assert (doc.ontology.revision, doc.dirty, len(doc.undo)) == before


def model_state(pid: str) -> tuple:
    doc = model_doc(pid)
    return doc.ontology.revision, doc.dirty, len(doc.undo)


# --- D1: the wizard, end to end ---------------------------------------------------


def test_d1_people_imported_with_exact_values_and_a_clean_report(pid):
    inspected = client.post(f"/api/projects/{pid}/data/inspect", files=csv_file(PEOPLE)).json()
    assert inspected["idSuggestion"] == "id" and inspected["nameSuggestion"] == "name"
    step2 = client.post(f"/api/projects/{pid}/data/preview", files=csv_file(PEOPLE),
                        data={"choices": json.dumps({"classIri": PERSON, "idColumn": "id"})}).json()
    assert step2["idCheck"]["ok"] is True
    # Name and birth date suggested from the headers, never applied by the server.
    assert step2["suggestions"]["name"] == {"as": "name"}
    assert step2["suggestions"]["birth date"] == {"as": "attribute", "property": EX + "birthDate"}
    assert step2["suggestions"]["nickname"] == {"as": "attribute", "property": EX + "nickname"}
    assert "org" not in step2["suggestions"]
    assert {f["label"] for f in step2["fields"]} == {"birth date", "nickname", "member of"}

    before = model_state(pid)
    made = import_csv(pid)["snapshot"]
    untouched(pid, before)
    # Every value fits and every row has an id. Since Stage B (5.9) the org
    # column's links are counted too, and no Organization rows are imported
    # here, so the report is not clean on that one line alone.
    assert made["report"]["keptAsText"] == [] and made["report"]["skipped"]["count"] == 0
    assert made["report"]["unmatched"] == [{
        "column": "org", "className": "Organization", "classIri": EX + "Organization", "count": 2, "rows": [1, 2],
    }]
    assert made["report"]["clean"] is False
    assert made["report"]["individuals"] == 3 and made["report"]["rowsRead"] == 3
    assert made["report"]["empty"] == [{"column": "nickname", "count": 1}, {"column": "org", "count": 1}]
    alice = URIRef(EX + "data/person/1")
    g = data_graph(pid, made["id"])
    assert set(g.predicate_objects(alice)) == {
        (RDF.type, U("Person")),
        (RDFS.label, Literal("Alice Martin", lang="en")),
        (U("birthDate"), Literal("1990-01-02", datatype=XSD.date)),
        (U("nickname"), Literal("Al", lang="en")),
        (U("memberOf"), URIRef(EX + "data/organization/acme")),
    }
    assert len(g) == made["statements"] == 13
    # The folder holds what 5.6 lists, and the model file is untouched.
    assert sorted(p.name for p in folder(pid, made["id"]).iterdir()) == [
        "data.ttl", "mapping.rml.ttl", "snapshot.json", "source.csv"]
    assert "data/person" not in (project_store.folder(pid) / "model.ttl").read_text(encoding="utf-8")


def test_d1_the_preview_reads_as_rows_with_kept_text_marked(pid):
    text = PEOPLE.replace("1985-12-31", "yesterday")
    preview = client.post(f"/api/projects/{pid}/data/preview", files=csv_file(text),
                          data={"choices": json.dumps(choices())}).json()
    bob = preview["rows"][1]
    assert bob["name"] == "Bob Stone"
    assert bob["values"][0] == {"column": "birth date", "label": "birth date", "value": "yesterday",
                                "kind": "value", "datatype": "date", "fits": False}
    # An empty cell says nothing.
    assert [v["column"] for v in bob["values"]] == ["birth date", "org"]
    assert preview["report"]["statements"] == 13


def test_d2_semicolon_windows_1252_names_correct(pid):
    text = "id;name\n1;Hélène\n2;François\n"
    inspected = client.post(f"/api/projects/{pid}/data/inspect", files=csv_file(text, encoding="windows-1252")).json()
    assert (inspected["separator"], inspected["encoding"]) == (";", "windows-1252")
    made = import_csv(pid, text, {"classIri": PERSON, "idColumn": "id", "columns": {"name": {"as": "name"}}},
                      encoding="windows-1252")["snapshot"]
    g = data_graph(pid, made["id"])
    assert (URIRef(EX + "data/person/1"), RDFS.label, Literal("Hélène", lang="en")) in g
    assert (URIRef(EX + "data/person/2"), RDFS.label, Literal("François", lang="en")) in g
    mapping = (folder(pid, made["id"]) / "mapping.rml.ttl").read_text(encoding="utf-8")
    assert 'csvw:delimiter ";"' in mapping and 'csvw:encoding "windows-1252"' in mapping


def test_d3_blank_and_duplicate_ids_then_the_row_number(pid):
    text = "id,name\n1,Ann\n,Bob\n1,Cy\n2,Dee\n"
    step2 = client.post(f"/api/projects/{pid}/data/preview", files=csv_file(text),
                        data={"choices": json.dumps({"classIri": PERSON, "idColumn": "id"})}).json()
    check = step2["idCheck"]
    assert check["ok"] is False
    assert check["missing"] == {"count": 1, "rows": [2]}
    assert check["repeats"] == 1 and check["repeated"] == {"count": 1, "rows": [3]}
    made = import_csv(pid, text, {"classIri": PERSON, "idColumn": None, "columns": {"name": {"as": "name"}}})["snapshot"]
    g = data_graph(pid, made["id"])
    assert {str(s) for s in g.subjects(RDF.type, U("Person"))} == {EX + f"data/person/{n}" for n in (1, 2, 3, 4)}
    assert (URIRef(EX + "data/person/2"), RDFS.label, Literal("Bob", lang="en")) in g
    # The row number is a column of the copy, so the mapping stays standard RML.
    assert (folder(pid, made["id"]) / "source.csv").read_text().splitlines()[0] == "row,id,name"


def test_d4_bad_dates_kept_as_text_counted_listed_and_reported_by_validate(pid):
    rows = ["id,name,birth date"] + [
        f"{n},P{n},{'yesterday' if n % 2 == 0 else '1990-01-01'}" for n in range(1, 81)
    ]
    made = import_csv(pid, "\n".join(rows) + "\n",
                      {"classIri": PERSON, "idColumn": "id",
                       "columns": {"name": {"as": "name"}, "birth date": {"as": "attribute", "property": EX + "birthDate"}}}
                      )["snapshot"]
    kept = made["report"]["keptAsText"]
    assert kept[0]["column"] == "birth date" and kept[0]["count"] == 40 and kept[0]["datatype"] == "date"
    assert kept[0]["rows"] == list(range(2, 41, 2))  # the first twenty listed
    g = data_graph(pid, made["id"])
    assert (URIRef(EX + "data/person/2"), U("birthDate"), Literal("yesterday")) in g
    # A V-6 type rule reports each such row.
    shapes = (project_store.folder(pid) / "shapes.ttl").read_text(encoding="utf-8") + f"""
<{EX}BirthRules> a sh:NodeShape ; rdfs:label "Birth rules"@en ; sh:targetClass <{PERSON}> ;
    sh:property [ sh:path <{EX}birthDate> ; sh:datatype xsd:date ] .
"""
    shapes = "@prefix xsd: <http://www.w3.org/2001/XMLSchema#> .\n" + shapes
    assert client.put(f"/api/projects/{pid}/documents/shapes/source", json={"text": shapes}).status_code == 200
    result = validate(pid)
    panel = next(p for p in result["shapes"] if p["name"] == "Birth rules")
    assert panel["failingCount"] == 40 and panel["state"] == "fails"
    assert panel["data"] == [made["id"]]
    assert result["dataSources"][0]["id"] == made["id"]


def test_d5_a_new_attribute_from_whole_numbers_is_one_undo_step_and_typed(pid):
    inspected = client.post(f"/api/projects/{pid}/data/inspect", files=csv_file("id,age\n1,30\n2,41\n")).json()
    assert inspected["columns"][1]["wholeNumbers"] is True
    undo_before = len(model_doc(pid).undo)
    created = command(pid, "CreateDatatypeProperty", label="age", domain=PERSON, datatype=str(XSD.integer))
    assert len(model_doc(pid).undo) == undo_before + 1
    made = import_csv(pid, "id,age\n1,30\n2,41\n",
                      {"classIri": PERSON, "idColumn": "id", "columns": {"age": {"as": "attribute", "property": created["created"]}}}
                      )["snapshot"]
    assert (URIRef(EX + "data/person/1"), U("age"), Literal("30", datatype=XSD.integer)) in data_graph(pid, made["id"])


def test_d6_taxonomy_rows_become_concepts_in_the_scheme():
    project = _create("taxonomy-small")
    scheme = next(project_store_graph(project).subjects(RDF.type, SKOS.ConceptScheme))
    step2 = client.post(f"/api/projects/{project}/data/preview", files=csv_file("code,label,def\nA,Apple,A fruit\n"),
                        data={"choices": json.dumps({"classIri": str(SKOS.Concept), "idColumn": "code"})}).json()
    assert step2["suggestions"]["label"] == {"as": "name"}
    assert {f["label"] for f in step2["fields"]} == {"alternative name", "definition", "notation"}
    made = import_csv(project, "code,label,def\nA,Apple,A fruit\nB,Banana,\n", {
        "classIri": str(SKOS.Concept), "idColumn": "code",
        "columns": {"label": {"as": "name"}, "def": {"as": "attribute", "property": str(SKOS.definition)}},
    })["snapshot"]
    apple = URIRef(EX + "data/concept/A")
    g = data_graph(project, made["id"])
    assert (apple, RDF.type, SKOS.Concept) in g
    assert (apple, SKOS.prefLabel, Literal("Apple", lang="en")) in g
    assert (apple, SKOS.inScheme, scheme) in g and (apple, SKOS.topConceptOf, scheme) in g
    assert (apple, SKOS.definition, Literal("A fruit", lang="en")) in g
    concepts = client.get(f"/api/ontologies/{oid(project)}/hierarchy").json()["concepts"]["nodes"]
    assert concepts[str(apple)]["fromData"] == made["id"]


def project_store_graph(pid: str) -> Graph:
    return model_doc(pid).graph


def test_d7_refresh_with_the_same_headers_then_with_a_column_renamed(pid):
    made = import_csv(pid)["snapshot"]
    sid = made["id"]
    before = model_state(pid)
    newer = PEOPLE + "4,Dan Ray,1970-07-07,D,acme\n"
    refreshed = client.post(f"/api/projects/{pid}/data/{sid}/refresh", files=csv_file(newer, "people (2).csv")).json()
    assert refreshed["status"] == "imported"
    assert refreshed["snapshot"]["rows"] == 4 and refreshed["snapshot"]["source"] == "people (2).csv"
    assert URIRef(EX + "data/person/4") in set(data_graph(pid, sid).subjects())
    renamed = newer.replace("birth date", "born")
    mismatch = client.post(f"/api/projects/{pid}/data/{sid}/refresh", files=csv_file(renamed)).json()
    assert mismatch["status"] == "mismatch" and mismatch["missing"] == ["birth date"]
    assert mismatch["choices"]["columns"]["birth date"]["as"] == "attribute"
    # Nothing changed by the mismatch.
    assert len(data_graph(pid, sid)) == refreshed["snapshot"]["statements"]
    # The wizard's step 3 sends new choices for the new file.
    redone = client.post(
        f"/api/projects/{pid}/data/{sid}/refresh", files=csv_file(renamed),
        data={"choices": json.dumps(choices(**{"born": {"as": "attribute", "property": EX + "birthDate"}},
                                            **{"birth date": {"as": "ignore"}}) | {"columns": {
                                                "name": {"as": "name"}, "born": {"as": "attribute", "property": EX + "birthDate"}}})},
    ).json()
    assert redone["status"] == "imported"
    assert (URIRef(EX + "data/person/4"), U("birthDate"), Literal("1970-07-07", datatype=XSD.date)) in data_graph(pid, sid)
    untouched(pid, before)


def test_d8_switched_off_data_leaves_explore_query_and_validate_then_returns(pid):
    made = import_csv(pid)["snapshot"]
    sid, alice = made["id"], EX + "data/person/1"
    query = f"SELECT ?s WHERE {{ ?s a <{PERSON}> }}"

    def seen() -> dict:
        graph = client.get(f"/api/ontologies/{oid(pid)}/graph").json()
        tree = client.get(f"/api/ontologies/{oid(pid)}/hierarchy").json()
        rows = sparql(pid, query)["rows"]
        return {
            "graph": any(n["id"] == alice for n in graph["nodes"]),
            "tree": alice in tree.get("examples", {}).get("nodes", {}),
            "search": any(r["id"] == alice for r in client.get(
                f"/api/ontologies/{oid(pid)}/search", params={"q": "Alice Martin"}).json()),
            "query": any(r[0]["value"] == alice for r in rows),
            "validate": validate(pid)["dataSources"] != [],
        }

    assert set(seen().values()) == {True}
    before = model_state(pid)
    off = client.patch(f"/api/projects/{pid}/data/{sid}", json={"enabled": False}).json()
    assert off["snapshot"]["enabled"] is False
    assert set(seen().values()) == {False}
    assert client.patch(f"/api/projects/{pid}/data/{sid}", json={"enabled": True}).status_code == 200
    assert set(seen().values()) == {True}
    untouched(pid, before)
    # And with the imports switch on, the data shows too.
    graph = client.get(f"/api/ontologies/{oid(pid)}/graph", params={"imports": "true"}).json()
    assert any(n["id"] == alice and n["fromData"] == "people.csv" for n in graph["nodes"])


def test_d9_edit_as_rml_reruns_then_a_function_is_refused_by_name_with_the_data_kept(pid):
    made = import_csv(pid)["snapshot"]
    sid = made["id"]
    text = client.get(f"/api/projects/{pid}/data").json()["snapshots"][0]["mappingText"]
    changed = text.replace('rml:reference "nickname"', 'rml:reference "name"')
    assert changed != text
    edited = client.patch(f"/api/projects/{pid}/data/{sid}", json={"mapping": changed}).json()["snapshot"]
    assert edited["mapping"]["status"] == "ok"
    assert (URIRef(EX + "data/person/2"), U("nickname"), Literal("Bob Stone", lang="en")) in data_graph(pid, sid)
    before = data_graph(pid, sid)
    function = changed.replace("rml:subjectMap [", "rml:subjectMap [ rml:function <http://example.org/f> ;")
    refused = client.patch(f"/api/projects/{pid}/data/{sid}", json={"mapping": function}).json()["snapshot"]
    assert refused["mapping"]["status"] == "outside"
    assert refused["mapping"]["message"].startswith("rml:function is outside what Semantic Studio runs")
    # The last good data stays, and the mapping is kept as written.
    assert len(data_graph(pid, sid)) == len(before)
    assert "rml:function" in (folder(pid, sid) / "mapping.rml.ttl").read_text(encoding="utf-8")


@pytest.mark.parametrize("source", ["../model.ttl", "C:\\\\x.csv", "https://example.org/people.csv"])
def test_d10_an_edited_source_naming_another_file_or_a_url_is_refused(pid, source):
    made = import_csv(pid)["snapshot"]
    sid = made["id"]
    text = (folder(pid, sid) / "mapping.rml.ttl").read_text(encoding="utf-8")
    before = (folder(pid, sid) / "data.ttl").read_bytes()
    pointed = text.replace('csvw:url "source.csv"', f'csvw:url "{source}"')
    assert pointed != text
    response = client.patch(f"/api/projects/{pid}/data/{sid}", json={"mapping": pointed})
    assert response.status_code == 422
    assert "must name this snapshot's own source.csv" in response.json()["detail"]
    # Nothing kept: not the mapping, not the data.
    assert (folder(pid, sid) / "mapping.rml.ttl").read_text(encoding="utf-8") == text
    assert (folder(pid, sid) / "data.ttl").read_bytes() == before


def test_d10_a_source_hidden_beside_an_unsupported_feature_is_still_refused(pid):
    made = import_csv(pid)["snapshot"]
    sid = made["id"]
    text = (folder(pid, sid) / "mapping.rml.ttl").read_text(encoding="utf-8")
    both = text.replace('csvw:url "source.csv"', 'csvw:url "../model.ttl"').replace(
        "rml:subjectMap [", "rml:subjectMap [ rml:function <http://example.org/f> ;")
    assert client.patch(f"/api/projects/{pid}/data/{sid}", json={"mapping": both}).status_code == 422
    assert (folder(pid, sid) / "mapping.rml.ttl").read_text(encoding="utf-8") == text


# --- D11: the 2,000-row limit on every path, and the label on a sample ---------------------


def people_rows(n: int) -> str:
    return "id,name\n" + "".join(f"{i},Person {i}\n" for i in range(1, n + 1))


NAME_ONLY = {"classIri": PERSON, "idColumn": "id", "columns": {"name": {"as": "name"}}}


def test_d11_exactly_two_thousand_rows_import(pid):
    made = import_csv(pid, people_rows(2000), NAME_ONLY)["snapshot"]
    assert made["rows"] == 2000 and made["total"] == 2000 and made["sample"] is False


def test_d11_one_row_more_needs_the_sample_choice(pid):
    response = client.post(f"/api/projects/{pid}/data", files=csv_file(people_rows(2001)),
                           data={"choices": json.dumps(NAME_ONLY)})
    assert response.status_code == 422
    assert response.json()["detail"].startswith(
        "This file has 2,001 rows. Semantic Studio imports at most 2,000 rows: it is for learning how "
        "data is mapped to a model, not for loading whole datasets.")
    assert not (project_store.folder(pid) / "data").exists() or not any((project_store.folder(pid) / "data").iterdir())


def test_d11_a_sample_holds_rows_one_to_two_thousand_and_says_so_everywhere(pid):
    made = import_csv(pid, people_rows(12_480), NAME_ONLY, options={"sample": True})["snapshot"]
    sid = made["id"]
    assert (made["rows"], made["total"], made["sample"]) == (2000, 12_480, True)
    assert made["report"]["total"] == 12_480 and made["report"]["sample"] is True
    subjects = {str(s) for s in data_graph(pid, sid).subjects(RDF.type, U("Person"))}
    assert EX + "data/person/1" in subjects and EX + "data/person/2000" in subjects
    assert EX + "data/person/2001" not in subjects and len(subjects) == 2000
    # The copy holds the rows imported, so another engine gives the same.
    assert len((folder(pid, sid) / "source.csv").read_text().splitlines()) == 2001
    meta = json.loads((folder(pid, sid) / "snapshot.json").read_text(encoding="utf-8"))
    assert meta["rows"] == {"kept": 2000, "total": 12_480}

    def sample(info: dict) -> bool:
        return (info["rows"], info["total"], info["sample"]) == (2000, 12_480, True)

    # Each place 5.6 lists receives both counts: the list and the card,
    # the detail panel, the tree, the Query source line and Validate.
    assert sample(client.get(f"/api/projects/{pid}/data").json()["snapshots"][0])
    card = next(p for p in client.get("/api/projects").json() if p["id"] == pid)
    assert sample(card["data"][0])
    detail = node(pid, EX + "data/person/7")["fromData"]
    assert sample(detail) and detail["row"] == 7 and detail["source"] == "people.csv"
    tree = client.get(f"/api/ontologies/{oid(pid)}/hierarchy").json()["examples"]["nodes"]
    assert tree[EX + "data/person/7"]["fromData"] == sid
    assert sample(sparql(pid, "SELECT ?s WHERE { ?s ?p ?o } LIMIT 1")["dataSources"][0])
    result = validate(pid)
    assert sample(result["dataSources"][0])
    assert any(sid in p["data"] for p in result["shapes"])


def test_d11_refresh_with_a_larger_file_stops_at_two_thousand(pid):
    sid = import_csv(pid, people_rows(10), NAME_ONLY)["snapshot"]["id"]
    refused = client.post(f"/api/projects/{pid}/data/{sid}/refresh", files=csv_file(people_rows(5000)))
    assert refused.status_code == 422 and "5,000 rows" in refused.json()["detail"]
    done = client.post(f"/api/projects/{pid}/data/{sid}/refresh", files=csv_file(people_rows(5000)),
                       data={"options": json.dumps({"sample": True})}).json()
    assert (done["snapshot"]["rows"], done["snapshot"]["total"]) == (2000, 5000)
    assert len(set(data_graph(pid, sid).subjects(RDF.type, U("Person")))) == 2000


def test_d11_an_edited_mapping_runs_on_two_thousand_rows_at_most(pid, monkeypatch):
    sid = import_csv(pid, people_rows(12_480), NAME_ONLY, options={"sample": True})["snapshot"]["id"]
    # Even a copy that somehow held more rows: the engine stops at 2,000.
    source = folder(pid, sid) / "source.csv"
    source.write_text(people_rows(3000), encoding="utf-8")
    text = (folder(pid, sid) / "mapping.rml.ttl").read_text(encoding="utf-8")
    edited = client.patch(f"/api/projects/{pid}/data/{sid}", json={"mapping": text + "\n"}).json()["snapshot"]
    assert edited["report"]["rowsRead"] == 2000
    subjects = {str(s) for s in data_graph(pid, sid).subjects(RDF.type, U("Person"))}
    assert len(subjects) == 2000 and EX + "data/person/2001" not in subjects


def test_d11_the_files_refused_with_their_sentences(pid):
    big = b"id,text\n" + b"1," + b"x" * (5 * 1024 * 1024) + b"\n"
    response = client.post(f"/api/projects/{pid}/data/inspect", files={"file": ("big.csv", big, "text/csv")})
    assert response.status_code == 413
    assert "Semantic Studio imports files of at most 5 MB" in response.json()["detail"]["message"]
    wide = ",".join(f"c{i}" for i in range(101)) + "\n"
    response = client.post(f"/api/projects/{pid}/data/inspect", files=csv_file(wide))
    assert response.status_code == 413
    assert response.json()["detail"]["message"].startswith("This file has 101 columns.")
    long_cell = "id,text\n1," + "x" * 40_000 + "\n"
    response = client.post(f"/api/projects/{pid}/data/inspect", files=csv_file(long_cell))
    assert response.status_code == 413 and "40,000 characters" in response.json()["detail"]["message"]


def test_d11_five_megabytes_and_one_byte_refused_while_reading(pid):
    data = b"id\n" + b"1\n" * ((5 * 1024 * 1024 - 3) // 2) + b"1"
    data = data + b"\n" * (5 * 1024 * 1024 - len(data))
    assert len(data) == 5 * 1024 * 1024
    assert client.post(f"/api/projects/{pid}/data/inspect", files={"file": ("a.csv", data, "text/csv")}).status_code == 200
    response = client.post(f"/api/projects/{pid}/data/inspect", files={"file": ("a.csv", data + b"1", "text/csv")})
    assert response.status_code == 413


def test_a_declared_oversize_is_refused_before_the_body_is_read(pid):
    response = client.post(
        f"/api/projects/{pid}/data", content=b"x",
        headers={"Content-Type": "multipart/form-data; boundary=x", "Content-Length": str(6 * 1024 * 1024)},
    )
    assert response.status_code == 413


# --- D12: the project zip, opened elsewhere ---------------------------------------------


def test_d12_a_zipped_project_opens_elsewhere_with_its_data_and_runs_nothing(pid, monkeypatch):
    made = import_csv(pid)["snapshot"]
    assert client.post(f"/api/projects/{pid}/documents/model/save").status_code == 200
    data = client.get(f"/api/projects/{pid}/export").content
    names = zipfile.ZipFile(io.BytesIO(data)).namelist()
    assert f"data/{made['id']}/data.ttl" in names and f"data/{made['id']}/snapshot.json" in names
    elsewhere = Path(tempfile.mkdtemp(prefix="semantic-studio-elsewhere-"))
    try:
        projects = ProjectStore(elsewhere)
        zipfile.ZipFile(io.BytesIO(data)).extractall(projects.dir / pid)
        service = EditingService(projects, OntologyStore(elsewhere))

        def never(*args, **kwargs):
            raise AssertionError("the mapping was run on open")

        monkeypatch.setattr(rml, "run", never)
        service.open(pid)
        generation, active = service.snapshots.active(pid)
        assert len(active) == 1 and len(active[0][0]) == made["statements"]
        assert active[0][1]["source"] == "people.csv"
        service.close_all()
    finally:
        shutil.rmtree(elsewhere, ignore_errors=True)


# --- the snapshot's lifecycle and refusals ------------------------------------------------


def test_remove_moves_the_folder_to_the_projects_trash_and_says_how_many_statements(pid):
    made = import_csv(pid)["snapshot"]
    before = model_state(pid)
    removed = client.delete(f"/api/projects/{pid}/data/{made['id']}").json()
    assert removed["statements"] == 13 and removed["location"] == f".trash/data-{made['id']}"
    assert (project_store.folder(pid) / ".trash" / f"data-{made['id']}" / "data.ttl").is_file()
    assert client.get(f"/api/projects/{pid}/data").json()["snapshots"] == []
    assert not any(n.get("fromData") for n in client.get(f"/api/ontologies/{oid(pid)}/graph").json()["nodes"])
    untouched(pid, before)
    # The project's own trash does not travel in its zip.
    names = zipfile.ZipFile(io.BytesIO(client.get(f"/api/projects/{pid}/export").content)).namelist()
    assert not any(n.startswith(".trash") for n in names)


def test_data_ttl_is_read_as_rdflib_reads_it_and_a_hand_edit_falls_back(tmp_path):
    from rdflib.compare import isomorphic

    from app.snapshots import SnapshotService, read_ntriples

    g = Graph()
    s = URIRef(EX + "data/person/1")
    for value in ['a "quoted" word', "back\\slash", "two\nlines\ttab", "é and 😀", ""]:
        g.add((s, U("text"), Literal(value, lang="en-GB")))
    g.add((s, U("n"), Literal("42", datatype=XSD.integer)))
    g.add((s, U("plain"), Literal("plain")))
    g.add((s, RDF.type, U("Person")))
    text = g.serialize(format="nt", encoding="utf-8").decode("utf-8")
    assert isomorphic(read_ntriples(text), g)
    with pytest.raises(ValueError):
        read_ntriples('_:b <http://x/p> "blank nodes are not ours" .')
    edited = tmp_path / "data.ttl"
    edited.write_text(f"@prefix ex: <{EX}> .\nex:a ex:b [ ex:c 1 ] .\n", encoding="utf-8")
    assert len(SnapshotService._parse_data(edited)) == 2


def test_a_switch_or_remove_while_closed_keeps_the_cards_line_true(pid):
    sid = import_csv(pid)["snapshot"]["id"]
    import_csv(pid, name="more.csv")
    editing_service.close(pid, discard=True)
    assert client.patch(f"/api/projects/{pid}/data/{sid}", json={"enabled": False}).status_code == 200
    card = next(p for p in client.get("/api/projects").json() if p["id"] == pid)["data"]
    assert sorted((d["source"], d["enabled"]) for d in card) == [("more.csv", True), ("people.csv", False)]
    assert client.delete(f"/api/projects/{pid}/data/{sid}").status_code == 200
    card = next(p for p in client.get("/api/projects").json() if p["id"] == pid)["data"]
    assert [d["source"] for d in card] == ["more.csv"]


def test_change_the_mapping_reads_the_copy_kept(pid):
    made = import_csv(pid)["snapshot"]
    sid = made["id"]
    stored = client.post(f"/api/projects/{pid}/data/inspect", data={"snapshot": sid}).json()
    assert [c["name"] for c in stored["columns"]] == ["id", "name", "birth date", "nickname", "org"]
    assert stored["choices"]["idColumn"] == "id"
    remapped = client.patch(f"/api/projects/{pid}/data/{sid}", json={"choices": NAME_ONLY}).json()["snapshot"]
    assert remapped["statements"] == 6
    assert set(data_graph(pid, sid).predicates()) == {RDF.type, RDFS.label}


def test_a_snapshot_individual_is_read_only_with_its_source_and_row(pid):
    import_csv(pid)
    details = node(pid, EX + "data/person/2")
    assert details["fromData"]["source"] == "people.csv" and details["fromData"]["row"] == 2
    # Never an example's form: snapshot data is not edited here (D-073).
    assert "example" not in details
    # A model example keeps its form.
    assert "example" in node(pid, EX + "alice")


@pytest.mark.parametrize("sid", ["../x", "people", "PEOPLE-abcdef", "people-abcdeg", "a" * 50 + "-abcdef"])
def test_a_snapshot_id_not_issued_is_404(pid, sid):
    assert client.patch(f"/api/projects/{pid}/data/{sid}", json={"enabled": False}).status_code in (404, 405)
    assert client.delete(f"/api/projects/{pid}/data/{sid}").status_code in (404, 405)


def test_the_wizard_needs_the_project_open(pid):
    editing_service.close(pid, discard=True)
    response = client.post(f"/api/projects/{pid}/data/preview", files=csv_file(PEOPLE),
                           data={"choices": json.dumps(choices())})
    assert response.status_code == 409


def test_the_routes_refuse_a_request_without_the_client_header(pid):
    bare = TestClient(app, base_url="http://localhost")
    assert bare.post(f"/api/projects/{pid}/data", files=csv_file(PEOPLE)).status_code in (400, 403)


def test_a_column_choice_outside_the_classs_fields_is_refused(pid):
    bad = choices(name={"as": "attribute", "property": EX + "memberOf"})
    response = client.post(f"/api/projects/{pid}/data", files=csv_file(PEOPLE), data={"choices": json.dumps(bad)})
    assert response.status_code == 422 and "Column name" in response.json()["detail"]


# --- budgets (Section 10) -------------------------------------------------------------------


def _full(pid: str, name: str) -> None:
    """2,000 rows x 20 columns: a full snapshot."""
    attrs = [command(pid, "CreateDatatypeProperty", label=f"{name} field {i}", domain=PERSON)["created"]
             for i in range(18)]
    header = ["id", "name"] + [f"{name} field {i}" for i in range(18)]
    lines = [",".join(header)] + [
        ",".join([f"{name}{n}", f"Person {n}"] + [f"v{n}-{i}" for i in range(18)]) for n in range(1, 2001)
    ]
    columns = {"name": {"as": "name"}} | {f"{name} field {i}": {"as": "attribute", "property": attrs[i]} for i in range(18)}
    import_csv(pid, "\n".join(lines) + "\n", {"classIri": PERSON, "idColumn": "id", "columns": columns},
               name=f"{name}.csv")


def _median_ms(action, samples: int = 5) -> float:
    times = []
    gc.disable()
    try:
        for _ in range(samples):
            start = time.perf_counter()
            action()
            times.append((time.perf_counter() - start) * 1000)
    finally:
        gc.enable()
    return sorted(times)[samples // 2]


@pytest.mark.perf
def test_open_with_two_full_snapshots_budget(pid):
    _full(pid, "first")
    _full(pid, "second")
    assert client.post(f"/api/projects/{pid}/documents/model/save").status_code == 200

    def reopen():
        editing_service.close(pid, discard=True)
        client.post(f"/api/projects/{pid}/open")

    median = _median_ms(reopen)
    assert len(editing_service.snapshots.active(pid)[1]) == 2
    assert median <= limit_ms(1000), f"opening with two full snapshots took {median:.0f} ms (median of 5)"


@pytest.mark.perf
def test_validate_with_two_full_snapshots_budget(pid):
    _full(pid, "first")
    _full(pid, "second")
    rules = "\n".join(
        f"""<{EX}Timed{i}> a sh:NodeShape ; rdfs:label "Timed {i}"@en ; sh:targetClass <{PERSON}> ;
            sh:property [ sh:path {path} ; {rule} ] ."""
        for i, (path, rule) in enumerate([
            ("rdfs:label", "sh:minCount 1"),
            ("rdfs:label", "sh:maxCount 1"),
            (f"<{EX}firstField0>", "sh:datatype rdf:langString"),
            (f"<{EX}secondField1>", "sh:minLength 2"),
            (f"<{EX}memberOf>", f"sh:class <{EX}Organization>"),
        ])
    )
    shapes = ("@prefix sh: <http://www.w3.org/ns/shacl#> .\n@prefix rdfs: <http://www.w3.org/2000/01/rdf-schema#> .\n"
              "@prefix rdf: <http://www.w3.org/1999/02/22-rdf-syntax-ns#> .\n" + rules)
    assert client.put(f"/api/projects/{pid}/documents/shapes/source", json={"text": shapes}).status_code == 200
    median = _median_ms(lambda: validate(pid))
    assert median <= limit_ms(2000), f"validating two full snapshots with 5 shapes took {median:.0f} ms (median of 5)"


# --- Stage B: Excel (5.8, row X1) ------------------------------------------------------------

ORGANIZATION = EX + "Organization"
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
ORG_IDS = ["acme", "beta", "gamma", "delta", "epsilon"]


def xlsx(build) -> bytes:
    from openpyxl import Workbook

    book = Workbook()
    build(book)
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


def orgs_workbook(extra: int = 0) -> bytes:
    """X1's orgs.xlsx: a notes sheet first, then Orgs with a title row and a
    blank row above the table, dates and numbers."""
    def build(book):
        notes = book.active
        notes.title = "Notes"
        notes.append(["Read me first"])
        sheet = book.create_sheet("Orgs")
        sheet.append(["Organizations, October 2026"])
        sheet.append([])
        sheet.append(["id", "name", "founded", "staff"])
        for n, oid_ in enumerate(ORG_IDS + [f"extra{i}" for i in range(extra)], start=1):
            sheet.append([oid_, oid_.title(), datetime.date(1990 + n, n, n), n * 10])
    return xlsx(build)


def org_choices(**columns) -> dict:
    base = {
        "name": {"as": "name"},
        "founded": {"as": "attribute", "property": EX + "founded"},
        "staff": {"as": "attribute", "property": EX + "staff"},
    }
    base.update(columns)
    return {"classIri": ORGANIZATION, "idColumn": "id", "columns": base}


def with_org_attributes(pid: str) -> None:
    command(pid, "CreateDatatypeProperty", label="founded", domain=ORGANIZATION, datatype=str(XSD.date))
    command(pid, "CreateDatatypeProperty", label="staff", domain=ORGANIZATION, datatype=str(XSD.integer))


def import_orgs(pid: str, data: bytes | None = None, options: dict | None = None, status: int = 200) -> dict:
    response = client.post(
        f"/api/projects/{pid}/data",
        files={"file": ("orgs.xlsx", data or orgs_workbook(), XLSX)},
        data={"choices": json.dumps(org_choices()),
              "options": json.dumps(options if options is not None else {"sheet": "Orgs", "headerRow": 3})},
    )
    assert response.status_code == status, response.text
    return response.json()


def test_a_workbook_with_an_empty_first_sheet_and_row_one_inspects_and_keeps_its_pickers(pid):
    # PR #54 review: row 1 empty or the first sheet empty answered 422 with
    # no sheet list, so step 1 had no picker to change either.
    def build(book):
        book.active.title = "Cover"
        orgs = book.create_sheet("Orgs")
        orgs["A3"], orgs["B3"] = "id", "name"
        orgs["A4"], orgs["B4"] = "acme", "Acme"

    data = xlsx(build)

    def inspect(options: dict):
        return client.post(
            f"/api/projects/{pid}/data/inspect",
            files={"file": ("orgs.xlsx", data, XLSX)}, data={"options": json.dumps(options)},
        )

    found = inspect({})
    assert found.status_code == 200, found.text
    shown = found.json()
    assert (shown["workbook"]["sheet"], shown["workbook"]["headerRow"]) == ("Orgs", 3)
    assert [c["name"] for c in shown["columns"]] == ["id", "name"] and shown["total"] == 1
    for options, sentence in (
        ({"sheet": "Orgs", "headerRow": 1}, "Row 1 of sheet Orgs is empty"),
        ({"sheet": "Cover"}, "Sheet Cover is empty"),
    ):
        refused = inspect(options)
        assert refused.status_code == 422
        detail = refused.json()["detail"]
        assert detail["message"].startswith(sentence) and detail["kind"] == "empty"
        assert [s["name"] for s in detail["workbook"]["sheets"]] == ["Cover", "Orgs"]
        assert detail["workbook"]["sheet"] == options["sheet"]


def test_x1_a_workbook_imports_by_sheet_and_header_row_with_types_kept_and_a_csv_copy(pid):
    with_org_attributes(pid)
    data = orgs_workbook()
    first = client.post(f"/api/projects/{pid}/data/inspect", files={"file": ("orgs.xlsx", data, XLSX)}).json()
    # The first sheet, header in row 1, until the learner chooses.
    assert first["format"] == "xlsx" and first["workbook"]["sheet"] == "Notes"
    # Each sheet counted below a header in row 1: Orgs' title row is its header there.
    assert first["workbook"]["sheets"] == [{"name": "Notes", "rows": 0}, {"name": "Orgs", "rows": 6}]
    chosen = client.post(f"/api/projects/{pid}/data/inspect", files={"file": ("orgs.xlsx", data, XLSX)},
                         data={"options": json.dumps({"sheet": "Orgs", "headerRow": 3})}).json()
    assert [c["name"] for c in chosen["columns"]] == ["id", "name", "founded", "staff"]
    assert chosen["total"] == 5 and chosen["idSuggestion"] == "id"
    assert chosen["workbook"]["top"][0] == {"row": 1, "cells": ["Organizations, October 2026"]}

    before = model_state(pid)
    made = import_orgs(pid, data)["snapshot"]
    untouched(pid, before)
    assert made["workbook"] == {"sheet": "Orgs", "headerRow": 3}
    assert made["report"]["clean"] is True and made["report"]["individuals"] == 5
    acme = URIRef(EX + "data/organization/acme")
    assert set(data_graph(pid, made["id"]).predicate_objects(acme)) == {
        (RDF.type, U("Organization")),
        (RDFS.label, Literal("Acme", lang="en")),
        (U("founded"), Literal("1991-01-01", datatype=XSD.date)),
        (U("staff"), Literal("10", datatype=XSD.integer)),
    }
    # The original kept, and the sheet written as UTF-8 CSV with a comma,
    # which is what the standard RML mapping reads.
    where = folder(pid, made["id"])
    assert (where / "source.xlsx").read_bytes() == data
    assert (where / "source.csv").read_bytes().startswith(
        b"id,name,founded,staff\r\nacme,Acme,1991-01-01,10\r\n")
    text = (where / "mapping.rml.ttl").read_text(encoding="utf-8")
    mapping = rml.read(text, EX)
    assert (mapping.separator, mapping.encoding) == (",", "utf-8")
    assert "xlsx" not in text
    zipped = zipfile.ZipFile(io.BytesIO(client.get(f"/api/projects/{pid}/export").content)).namelist()
    assert f"data/{made['id']}/source.xlsx" in zipped and f"data/{made['id']}/source.csv" in zipped


def test_x1_the_row_limit_and_the_sample_apply_per_sheet(pid):
    def build(book):
        book.active.title = "Small"
        book.active.append(["id", "name"])
        book.active.append(["solo", "Solo"])
        big = book.create_sheet("Big")
        big.append(["id", "name"])
        for i in range(1, 5001):
            big.append([f"o{i}", f"Org {i}"])

    data = xlsx(build)
    found = client.post(f"/api/projects/{pid}/data/inspect", files={"file": ("big.xlsx", data, XLSX)},
                        data={"options": json.dumps({"sheet": "Big"})}).json()
    assert found["total"] == 5000 and found["sample"] is True
    assert found["limitSentence"].startswith("This file has 5,000 rows. Semantic Studio imports at most 2,000 rows")
    chosen = {"classIri": ORGANIZATION, "idColumn": "id", "columns": {"name": {"as": "name"}}}
    send = lambda options, status: client.post(  # noqa: E731
        f"/api/projects/{pid}/data", files={"file": ("big.xlsx", data, XLSX)},
        data={"choices": json.dumps(chosen), "options": json.dumps(options)})
    assert send({"sheet": "Big"}, 422).status_code == 422
    made = send({"sheet": "Big", "sample": True}, 200).json()["snapshot"]
    assert (made["rows"], made["total"], made["sample"]) == (2000, 5000, True)
    g = data_graph(pid, made["id"])
    assert (URIRef(EX + "data/organization/o2000"), RDF.type, U("Organization")) in g
    assert (URIRef(EX + "data/organization/o2001"), None, None) not in g
    # The other sheet of the same workbook is under the limit.
    small = send({"sheet": "Small"}, 200).json()["snapshot"]
    assert (small["rows"], small["total"], small["sample"]) == (1, 1, False)


def test_x1_refresh_reads_the_new_workbook_the_same_way_and_a_csv_drops_the_original(pid):
    with_org_attributes(pid)
    made = import_orgs(pid)["snapshot"]
    sid = made["id"]
    newer = orgs_workbook(extra=2)
    refreshed = client.post(f"/api/projects/{pid}/data/{sid}/refresh",
                            files={"file": ("orgs (2).xlsx", newer, XLSX)}).json()
    # No sheet or header row sent: the same ones as before.
    assert refreshed["status"] == "imported", refreshed
    assert refreshed["snapshot"]["rows"] == 7 and refreshed["snapshot"]["workbook"] == {"sheet": "Orgs", "headerRow": 3}
    assert (folder(pid, sid) / "source.xlsx").read_bytes() == newer
    # Another sheet named in the options starts from its own row 1, never
    # from the last sheet's header row (code review).
    def plain(book):
        book.active.title = "Plain"
        book.active.append(["id", "name", "founded", "staff"])
        book.active.append(["zeta", "Zeta", datetime.date(2000, 1, 1), 1])

    moved = client.post(f"/api/projects/{pid}/data/{sid}/refresh",
                        files={"file": ("plain.xlsx", xlsx(plain), XLSX)},
                        data={"options": json.dumps({"sheet": "Plain"})}).json()
    assert moved["status"] == "imported", moved
    assert moved["snapshot"]["workbook"] == {"sheet": "Plain", "headerRow": 1} and moved["snapshot"]["rows"] == 1
    as_csv = "id,name,founded,staff\nacme,Acme,1991-01-01,10\n"
    again = client.post(f"/api/projects/{pid}/data/{sid}/refresh", files=csv_file(as_csv, "orgs.csv")).json()
    assert again["status"] == "imported" and again["snapshot"]["workbook"] is None
    assert not (folder(pid, sid) / "source.xlsx").exists()
    assert (folder(pid, sid) / "source.csv").is_file()


# --- Stage B: links between snapshots (5.9, row X3) ---------------------------------------


# Twenty people: six link to an organization of orgs.xlsx, twelve to ones it
# does not hold, two have none.
LINKED_PEOPLE = "id,name,org\n" + "".join(
    f"{n},Person {n},{ORG_IDS[n % 5] if n <= 6 else ('' if n > 18 else f'gone{n}')}\n" for n in range(1, 21)
)
UNMATCHED_ROWS = list(range(7, 19))


def import_linked_people(pid: str) -> dict:
    return import_csv(pid, LINKED_PEOPLE, {
        "classIri": PERSON, "idColumn": "id",
        "columns": {"name": {"as": "name"}, "org": {"as": "relationship", "property": EX + "memberOf"}},
    })["snapshot"]


def people_row(pid: str) -> dict:
    return next(s for s in client.get(f"/api/projects/{pid}/data").json()["snapshots"] if s["source"] == "people.csv")


def members(pid: str) -> int:
    rows = sparql(pid, f"""SELECT (COUNT(?p) AS ?n) WHERE {{
        ?p <{EX}memberOf> ?o . ?o a <{ORGANIZATION}> . FILTER(STRSTARTS(STR(?p), "{EX}data/person/")) }}""")["rows"]
    return int(rows[0][0]["value"])


def test_x3_people_first_then_orgs_links_resolve_and_twelve_are_reported(pid):
    with_org_attributes(pid)
    people = import_linked_people(pid)
    # No Organization rows yet: every link matches none, and says so.
    assert people["report"]["unmatched"][0]["count"] == 18
    assert people["report"]["clean"] is False
    import_orgs(pid)
    row = people_row(pid)
    assert row["report"]["unmatched"] == [{
        "column": "org", "className": "Organization", "classIri": ORGANIZATION,
        "count": 12, "rows": UNMATCHED_ROWS,
    }]
    assert members(pid) == 6
    # The twelve are still written, so the Small template's Points to rule
    # (Person rules: member of points to an Organization) reports them.
    gone = {URIRef(f"{EX}data/organization/gone{n}") for n in UNMATCHED_ROWS}
    assert {o for o in data_graph(pid, people["id"]).objects(None, U("memberOf"))} >= gone
    panel = next(p for p in validate(pid)["shapes"] if p["name"] == "Person rules")
    # A problem's value is shown in short form: data/organization/gone7.
    flagged = {(p["focus"], p["value"]) for p in panel["problems"] if p["value"] and "gone" in p["value"]}
    assert flagged == {(f"{EX}data/person/{n}", f"data/organization/gone{n}") for n in UNMATCHED_ROWS}
    assert people["id"] in panel["data"]
    # Counted when asked: the orgs switched off, every link is to no row again.
    orgs = next(s for s in client.get(f"/api/projects/{pid}/data").json()["snapshots"] if s["source"] == "orgs.xlsx")
    assert client.patch(f"/api/projects/{pid}/data/{orgs['id']}", json={"enabled": False}).status_code == 200
    assert people_row(pid)["report"]["unmatched"][0]["count"] == 18


def test_x3_orgs_first_gives_the_same_statements_and_the_same_report(pid):
    with_org_attributes(pid)
    other = _create()
    with_org_attributes(other)
    # Order A, then order B, in two projects on the same base IRI.
    a_people = import_linked_people(pid)
    import_orgs(pid)
    import_orgs(other)
    preview = client.post(f"/api/projects/{other}/data/preview", files=csv_file(LINKED_PEOPLE), data={
        "choices": json.dumps({"classIri": PERSON, "idColumn": "id",
                               "columns": {"name": {"as": "name"},
                                           "org": {"as": "relationship", "property": EX + "memberOf"}}})}).json()
    # Step 4 already says it, before the import.
    assert preview["report"]["unmatched"][0]["count"] == 12
    b_people = import_linked_people(other)
    assert b_people["report"]["unmatched"][0]["rows"] == UNMATCHED_ROWS
    assert set(data_graph(pid, a_people["id"])) == set(data_graph(other, b_people["id"]))
    assert people_row(pid)["report"] == people_row(other)["report"]
    assert members(pid) == members(other) == 6


def test_x3_a_link_uses_the_iri_pattern_of_the_target_classs_snapshot(pid):
    # An expert's mapping names organizations under data/company/; a column
    # of ids linking to Organization then uses that pattern, not the default.
    with_org_attributes(pid)
    orgs = import_orgs(pid)["snapshot"]
    text = (folder(pid, orgs["id"]) / "mapping.rml.ttl").read_text(encoding="utf-8")
    edited = text.replace(f"{EX}data/organization/", f"{EX}data/company/")
    assert edited != text
    applied = client.patch(f"/api/projects/{pid}/data/{orgs['id']}", json={"mapping": edited})
    assert applied.status_code == 200, applied.text
    people = import_linked_people(pid)
    targets = set(data_graph(pid, people["id"]).objects(None, U("memberOf")))
    assert URIRef(f"{EX}data/company/acme") in targets
    assert not any("data/organization/" in str(t) for t in targets)
    assert people["report"]["unmatched"][0]["count"] == 12
    assert members(pid) == 6
