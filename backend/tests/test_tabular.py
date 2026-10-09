"""
================================================================================
FILE: backend/tests/test_tabular.py
================================================================================

SUMMARY
    Reading a CSV file for the data import (csv-data-import 5.2, Section 9):
    encoding, separator and header detection, header names made usable, the
    fixed limits, and what step 1 and step 2 suggest. Stage B (5.8, row X1):
    an Excel workbook's sheet and header row, its cell types kept as text
    the type rule reads, the same limits per sheet, and a far cell read
    without filling the gap to it. The ZIP and XML defences are in
    test_xlsx_defences.py.

BASIC IDEA
    Every file is generated in the test, never committed (rdf-fixture). The
    limits are asserted at their edges -- 2,000 rows kept and 2,001 not,
    5 MB and one byte more, 100 columns and 101, 32,767 characters and
    40,000 -- because an off-by-one there is exactly the defect a limit
    test exists to catch. The 2,000-row limit is asserted not to be a
    setting at all (D-098). Workbooks are made with openpyxl in the test,
    or written by hand where Excel could not save one.

INPUTS / INPUT SOURCES
    - app.tabular, on bytes built here.

EXPECTED OUTPUT
    - Pass/fail; two `perf` budgets (Section 10: inspect a 5 MB file in 1 s,
      read a 5,000-row sheet in 1 s).
================================================================================
"""

from __future__ import annotations

import datetime
import gc
import inspect as pyinspect
import io
import time
import zipfile

import pytest

from app import tabular
from tests.budget import limit_ms


def csv(rows: list[list[str]], sep: str = ",") -> str:
    return "\n".join(sep.join(r) for r in rows) + "\n"


def people(n: int) -> bytes:
    lines = [["id", "name", "born"]] + [[str(i), f"Person {i}", "2000-01-01"] for i in range(1, n + 1)]
    return csv(lines).encode()


# --- detection (5.2) -------------------------------------------------------------


def test_comma_utf8_with_header():
    table = tabular.read_table(b"id,name\n1,Ann\n2,Bob\n")
    assert (table.separator, table.encoding, table.header) == (",", "utf-8", True)
    assert table.columns == ["id", "name"]
    assert table.rows == [["1", "Ann"], ["2", "Bob"]]
    assert table.total == 2 and not table.sample


def test_semicolon_windows_1252_accented_names():
    # D2: what French Excel writes.
    data = "id;nom\n1;Hélène\n2;François\n".encode("windows-1252")
    table = tabular.read_table(data)
    assert table.separator == ";"
    assert table.encoding == "windows-1252"
    assert [r[1] for r in table.rows] == ["Hélène", "François"]


def test_tab_and_utf8_bom():
    data = "﻿id\tname\n1\tÅsa\n".encode("utf-8")
    table = tabular.read_table(data)
    assert (table.separator, table.encoding) == ("\t", "utf-8-sig")
    # The mark is not part of the first column's name.
    assert table.columns == ["id", "name"]


def test_a_semicolon_inside_quoted_commas_does_not_fool_the_separator():
    data = b'id,note\n1,"a; b; c"\n2,"d; e"\n'
    assert tabular.read_table(data).separator == ","


def test_choices_replace_the_detections():
    data = b"a;b\n1;2\n"
    table = tabular.read_table(data, separator=",", encoding="utf-8", header=False)
    assert table.columns == ["column 1"]
    assert table.rows == [["a;b"], ["1;2"]]
    assert table.total == 2
    # Nothing was named, so nothing is said to be renamed.
    assert table.renamed == []


def test_a_wrong_encoding_choice_is_said():
    with pytest.raises(tabular.TabularError, match="cannot be read as utf-8"):
        tabular.read_table("é".encode("windows-1252"), encoding="utf-8")


def test_binary_is_not_csv_text():
    with pytest.raises(tabular.TabularError, match="not CSV text") as caught:
        tabular.read_table(b"PK\x03\x04\x00\x00binary")
    assert caught.value.kind == "not-text"


def test_an_empty_file_is_said():
    with pytest.raises(tabular.TabularError, match="empty"):
        tabular.read_table(b"\n\n")


def test_values_are_trimmed_and_blank_lines_are_no_rows():
    table = tabular.read_table(b"id,name\n 1 ,  Ann  \n\n2,Bob\n")
    assert table.rows == [["1", "Ann"], ["2", "Bob"]]
    assert table.total == 2


# --- header names (5.2) --------------------------------------------------------------


def test_empty_and_duplicate_headers_are_named_and_said():
    table = tabular.read_table(b"id,name,,name,name\n1,a,b,c,d\n")
    assert table.columns == ["id", "name", "column 3", "name (2)", "name (3)"]
    assert table.renamed == [
        {"index": 2, "from": "", "to": "column 3"},
        {"index": 3, "from": "name", "to": "name (2)"},
        {"index": 4, "from": "name", "to": "name (3)"},
    ]


def test_a_row_wider_than_the_header_keeps_its_cells():
    table = tabular.read_table(b"id,name\n1,Ann,extra\n2,Bob\n")
    assert table.columns == ["id", "name", "column 3"]
    assert table.rows == [["1", "Ann", "extra"], ["2", "Bob", ""]]


# --- the limits (5.2, Section 9, D-098) ---------------------------------------------


def test_exactly_two_thousand_rows_are_all_kept():
    table = tabular.read_table(people(2000))
    assert len(table.rows) == 2000 and table.total == 2000 and not table.sample


def test_one_row_more_is_a_sample_of_the_first_two_thousand():
    table = tabular.read_table(people(2001))
    assert table.total == 2001 and table.sample
    assert len(table.rows) == 2000
    assert table.rows[-1][0] == "2000"


def test_twelve_thousand_rows_are_counted_and_the_sentence_says_so():
    table = tabular.read_table(people(12_480))
    assert table.total == 12_480 and len(table.rows) == 2000
    shown = tabular.inspect(table)
    assert shown["limitSentence"] == (
        "This file has 12,480 rows. Semantic Studio imports at most 2,000 rows: it is for "
        "learning how data is mapped to a model, not for loading whole datasets."
    )
    assert shown["sample"] is True and shown["kept"] == 2000 and shown["total"] == 12_480


def test_the_row_limit_is_not_a_setting(monkeypatch):
    # D-098: fixed. No environment variable moves it, and the module reads none.
    monkeypatch.setenv("SEMANTIC_STUDIO_MAX_DATA_ROWS", "50000")
    assert tabular.MAX_ROWS == 2000
    # The one environment line there is a write, switching openpyxl's lxml
    # off (Section 9, PR #54 review).
    source = pyinspect.getsource(tabular).replace('os.environ["OPENPYXL_LXML"] = "False"', "")
    assert "os.environ" not in source and "getenv(" not in source


def test_five_megabytes_are_read_and_one_byte_more_is_refused():
    header = b"id,text\n"
    line = b"1," + b"x" * 1000 + b"\n"
    body = header + line * ((tabular.MAX_BYTES - len(header)) // len(line))
    data = body + b"y" * (tabular.MAX_BYTES - len(body))
    assert len(data) == tabular.MAX_BYTES
    tabular.read_table(data)
    with pytest.raises(tabular.TabularError) as caught:
        tabular.read_table(data + b"z")
    assert caught.value.kind == "too-large"
    assert "at most 5 MB: it is for learning how data is mapped to a model" in str(caught.value)


def test_a_hundred_columns_are_read_and_a_hundred_and_one_refused():
    tabular.read_table((",".join(f"c{i}" for i in range(100)) + "\n").encode())
    with pytest.raises(tabular.TabularError, match="This file has 101 columns. Semantic Studio imports at most 100 columns"):
        tabular.read_table((",".join(f"c{i}" for i in range(101)) + "\n").encode())


def test_a_cell_at_the_excel_limit_is_read_and_forty_thousand_refused():
    tabular.read_table(b"id,text\n1," + b"x" * 32_767 + b"\n")
    with pytest.raises(tabular.TabularError, match="holds 40,000 characters. A cell can hold at most 32,767"):
        tabular.read_table(b"id,text\n1," + b"x" * 40_000 + b"\n")


def test_a_cell_past_the_csv_modules_own_limit_is_the_same_refusal():
    with pytest.raises(tabular.TabularError, match="A cell can hold at most 32,767"):
        tabular.read_table(b'id,text\n1,"' + b"x" * 200_000 + b'"\n')


# --- what steps 1 and 2 suggest (5.2, 5.3) --------------------------------------------


def test_the_identifier_and_name_are_suggested():
    table = tabular.read_table(b"person_id,Full-Name,Title\n1,Ann,Dr\n2,Bob,Mr\n")
    shown = tabular.inspect(table)
    assert shown["idSuggestion"] == "person_id"
    assert shown["nameSuggestion"] == "Title"  # name, label, title -- Full-Name is not "name"
    table = tabular.read_table(b"code,Name\n1,Ann\n")
    assert tabular.suggest_name(table) == "Name"


def test_an_identifier_with_gaps_or_repeats_is_not_suggested_and_is_reported():
    table = tabular.read_table(b"id,name\n1,a\n,b\n1,c\n2,d\n,e\n2,f\n3,g\n")
    assert tabular.suggest_id(table) is None
    check = tabular.id_check(table, "id")
    assert check["ok"] is False
    assert check["missing"] == {"count": 2, "rows": [2, 5]}
    assert check["repeats"] == 2  # "1" and "2" repeat
    assert check["repeated"] == {"count": 2, "rows": [3, 6]}


def test_column_profiles_suggest_a_new_attributes_type():
    table = tabular.read_table(b"n,d,t\n1,2026-01-02,x\n-3,2026-02-28,\n")
    n, d, t = tabular.inspect(table)["columns"]
    assert n["wholeNumbers"] and not n["dates"]
    assert d["dates"] and not d["wholeNumbers"]
    assert not t["wholeNumbers"] and not t["dates"] and t["empty"] == 1
    # A date that has the shape of one and is not: not all dates.
    assert not tabular.inspect(tabular.read_table(b"d\n2026-02-30\n"))["columns"][0]["dates"]


def test_the_copy_is_written_in_the_dialect_with_a_row_column():
    table = tabular.read_table("id;nom\n1;Hélène\n".encode("windows-1252"))
    data = tabular.write_csv(table, "windows-1252", ";", "row")
    assert data == "row;id;nom\r\n1;1;Hélène\r\n".encode("windows-1252")
    # A byte-order mark is never written.
    assert not tabular.write_csv(table, "utf-8-sig", ",").startswith(b"\xef\xbb\xbf")


@pytest.mark.perf
def test_inspect_budget():
    # Section 10: a 5 MB file -- detections, 20 rows and the row count -- in 1 s.
    header = b"id,name,born,city,note\n"
    line = b"123456,Somebody Withaname,1990-01-01,Somewhere,a note of some length here\n"
    data = header + line * ((tabular.MAX_BYTES - len(header)) // len(line))
    times = []
    gc.disable()
    try:
        for _ in range(5):
            start = time.perf_counter()
            tabular.inspect(tabular.read_table(data))
            times.append((time.perf_counter() - start) * 1000)
    finally:
        gc.enable()
    median = sorted(times)[2]
    assert median <= limit_ms(1000), f"inspecting a 5 MB file took {median:.0f} ms (median of 5)"


# --- Excel (Stage B, 5.8; row X1) ------------------------------------------------


def workbook(build) -> bytes:
    """A workbook made in the test with openpyxl, as Excel would save it."""
    from openpyxl import Workbook

    book = Workbook()
    build(book)
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


def orgs_book(book) -> None:
    # X1: two sheets, a title row above the table, dates and numbers.
    sheet = book.active
    sheet.title = "Orgs"
    sheet["A1"] = "Organizations, October 2026"
    sheet.merge_cells("A1:D1")
    sheet.append([])
    sheet.append(["id", "name", "founded", "staff", "share", "checked", "audited"])
    sheet.append(["acme", "Acme", datetime.date(1990, 1, 2), 12, 0.25, True, datetime.datetime(2026, 9, 1, 14, 30)])
    sheet.append(["beta", "Beta", datetime.date(2001, 5, 6), 3.0, 1e-05, False, datetime.datetime(2026, 9, 2)])
    sheet.append(["gamma", "  Gamma  ", None, 7, 1.5, None, None])
    sheet.row_dimensions[6].hidden = True
    people = book.create_sheet("People")
    people.append(["id", "name", "org"])
    for i in range(1, 6):
        people.append([i, f"Person {i}", "acme"])


def test_a_sheet_and_header_row_are_chosen_and_cell_types_kept():
    data = workbook(orgs_book)
    table = tabular.read_workbook(data, filename="orgs.xlsx", sheet="Orgs", header_row=3)
    assert table.columns == ["id", "name", "founded", "staff", "share", "checked", "audited"]
    assert table.rows == [
        # Numbers stay numbers, dates become xsd:date or, with a time, xsd:dateTime.
        ["acme", "Acme", "1990-01-02", "12", "0.25", "true", "2026-09-01T14:30:00"],
        ["beta", "Beta", "2001-05-06", "3", "0.00001", "false", "2026-09-02"],
        # Trimmed (5.5), and the hidden row is read.
        ["gamma", "Gamma", "", "7", "1.5", "", ""],
    ]
    assert (table.separator, table.encoding, table.header, table.total) == (",", "utf-8", True, 3)
    found = tabular.inspect(table)
    assert found["format"] == "xlsx"
    assert found["workbook"]["sheets"] == [{"name": "Orgs", "rows": 3}, {"name": "People", "rows": 5}]
    assert (found["workbook"]["sheet"], found["workbook"]["headerRow"]) == ("Orgs", 3)
    # The header row picker numbers rows as Excel does, the empty one too.
    assert found["workbook"]["top"][:3] == [
        {"row": 1, "cells": ["Organizations, October 2026"]},
        {"row": 2, "cells": []},
        {"row": 3, "cells": ["id", "name", "founded", "staff"]},
    ]
    profiles = {c["name"]: c for c in found["columns"]}
    assert profiles["founded"]["dates"] and profiles["staff"]["wholeNumbers"] and profiles["share"]["numbers"]
    assert found["idSuggestion"] == "id" and found["nameSuggestion"] == "name"


def test_the_first_sheet_and_row_one_by_default_and_the_copy_is_utf8_comma():
    data = workbook(orgs_book)
    table = tabular.read_file(data, "orgs.xlsx", {"sheet": "People"})
    assert table.columns == ["id", "name", "org"] and table.total == 5
    assert table.rows[0] == ["1", "Person 1", "acme"]
    assert tabular.write_csv(table, "utf-8", ",").startswith(b"id,name,org\r\n1,Person 1,acme\r\n")
    first = tabular.read_file(data, "orgs.xlsx", {})
    assert first.workbook["sheet"] == "Orgs" and first.workbook["headerRow"] == 1
    # Row 1 is the merged title: its top-left value is the one column name.
    assert first.columns[0] == "Organizations, October 2026"


def test_a_formula_gives_its_last_saved_value_never_evaluated():
    def build(book):
        sheet = book.active
        sheet.append(["id", "total"])
        sheet.append(["a", "=1+1"])

    data = workbook(build)
    # openpyxl saves no cached value for a formula, as a file never opened in
    # Excel has none: the cell reads empty, and the formula is not run.
    assert tabular.read_workbook(data).rows == [["a", ""]]
    # A cached value, as Excel writes it, is what is read.
    with zipfile.ZipFile(io.BytesIO(data)) as zf:
        parts = {n: zf.read(n) for n in zf.namelist()}
    sheet = parts["xl/worksheets/sheet1.xml"].decode()
    parts["xl/worksheets/sheet1.xml"] = sheet.replace("<f>1+1</f><v></v>", "<f>1+1</f><v>2</v>").replace(
        "<f>1+1</f><v/>", "<f>1+1</f><v>2</v>").encode()
    if b"<v>2</v>" not in parts["xl/worksheets/sheet1.xml"]:
        parts["xl/worksheets/sheet1.xml"] = sheet.replace("<f>1+1</f>", "<f>1+1</f><v>2</v>").encode()
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zf:
        for name, raw in parts.items():
            zf.writestr(name, raw)
    assert tabular.read_workbook(out.getvalue()).rows == [["a", "2"]]


def test_the_row_limit_holds_per_sheet():
    def build(book):
        small = book.active
        small.title = "Small"
        small.append(["id"])
        small.append(["only"])
        big = book.create_sheet("Big")
        big.append(["id", "n"])
        for i in range(1, 5001):
            big.append([f"r{i}", i])

    data = workbook(build)
    table = tabular.read_workbook(data, sheet="Big")
    assert table.total == 5000 and len(table.rows) == tabular.MAX_ROWS and table.sample
    assert table.rows[-1] == ["r2000", "2000"]
    assert tabular.inspect(table)["limitSentence"] == tabular.rows_sentence(5000)
    small = tabular.read_workbook(data, sheet="Small")
    assert small.total == 1 and not small.sample


def test_header_row_and_sheet_refusals():
    data = workbook(orgs_book)
    with pytest.raises(tabular.TabularError, match="no sheet named Nope"):
        tabular.read_workbook(data, sheet="Nope")
    with pytest.raises(tabular.TabularError, match="Row 2 of sheet Orgs is empty: choose the row"):
        tabular.read_workbook(data, sheet="Orgs", header_row=2)
    for bad in (0, -1, "3", True, tabular.MAX_HEADER_ROW + 1):
        with pytest.raises(tabular.TabularError, match="The header row is a row number"):
            tabular.read_workbook(data, header_row=bad)

    def empty(book):
        book.active.title = "Blank"

    with pytest.raises(tabular.TabularError, match="Sheet Blank is empty") as caught:
        tabular.read_workbook(workbook(empty))
    assert caught.value.kind == "empty"
    # Even a workbook of empty sheets names them, for the Sheet picker.
    assert caught.value.workbook["sheets"] == [{"name": "Blank", "rows": 0}]


def late_table(book) -> None:
    # A table starting at row 3, nothing above it (PR #54 review).
    sheet = book.active
    sheet.title = "Late"
    sheet["A3"], sheet["B3"] = "id", "name"
    sheet["A4"], sheet["B4"] = "acme", "Acme"
    sheet["A5"], sheet["B5"] = "beta", "Beta"


def test_not_chosen_the_header_row_is_the_first_row_holding_a_value():
    table = tabular.read_file(workbook(late_table), "late.xlsx", {"sheet": "Late"})
    assert table.columns == ["id", "name"] and table.total == 2
    assert table.workbook["headerRow"] == 3
    assert [entry["row"] for entry in table.workbook["top"]] == [1, 2, 3, 4, 5]
    assert tabular.read_file(workbook(late_table), "late.xlsx", {}).workbook["headerRow"] == 3


def test_not_chosen_the_sheet_is_the_first_holding_a_value():
    def behind_an_empty_sheet(book):
        book.active.title = "Cover"
        people = book.create_sheet("People")
        people.append(["id", "name"])
        people.append(["1", "Ann"])

    table = tabular.read_file(workbook(behind_an_empty_sheet), "people.xlsx", {})
    assert table.workbook["sheet"] == "People" and table.columns == ["id", "name"]
    assert table.workbook["sheets"] == [{"name": "Cover", "rows": 0}, {"name": "People", "rows": 1}]
    # Chosen, the empty sheet is refused with the sheets still named.
    with pytest.raises(tabular.TabularError, match="Sheet Cover is empty") as caught:
        tabular.read_file(workbook(behind_an_empty_sheet), "people.xlsx", {"sheet": "Cover"})
    assert [s["name"] for s in caught.value.workbook["sheets"]] == ["Cover", "People"]


def test_a_chosen_empty_header_row_is_refused_with_the_sheets_and_rows_for_the_pickers():
    with pytest.raises(tabular.TabularError, match="Row 1 of sheet Late is empty") as caught:
        tabular.read_file(workbook(late_table), "late.xlsx", {"sheet": "Late", "headerRow": 1})
    listed = caught.value.workbook
    assert caught.value.kind == "empty"
    assert listed["sheet"] == "Late" and listed["headerRow"] == 1
    assert listed["sheets"] == [{"name": "Late", "rows": 0}]
    assert listed["top"] == [
        {"row": 1, "cells": []}, {"row": 2, "cells": []}, {"row": 3, "cells": ["id", "name"]},
        {"row": 4, "cells": ["acme", "Acme"]}, {"row": 5, "cells": ["beta", "Beta"]},
    ]


def test_a_sheet_holds_the_same_column_and_cell_limits_as_csv():
    def wide(book):
        book.active.append([f"c{i}" for i in range(1, 102)])

    with pytest.raises(tabular.TabularError, match="value in column 101") as caught:
        tabular.read_workbook(workbook(wide))
    assert caught.value.kind == "too-large"

    def hundred(book):
        book.active.append([f"c{i}" for i in range(1, 101)])
        book.active.append(["x"] * 100)

    assert len(tabular.read_workbook(workbook(hundred)).columns) == 100

    # Excel cannot save a longer cell (openpyxl cuts one), so the file is
    # written by hand, as an attacker would.
    base = workbook(lambda book: book.active.append(["id"]))
    for size, refused in ((tabular.MAX_CELL, False), (tabular.MAX_CELL + 1, True)):
        long_cell = _with_sheet_xml(
            base, '<row r="1"><c r="A1" t="inlineStr"><is><t>note</t></is></c></row>'
                  f'<row r="2"><c r="A2" t="inlineStr"><is><t>{"y" * size}</t></is></c></row>', "A1:A2")
        if refused:
            with pytest.raises(tabular.TabularError, match="holds 32,768 characters") as caught:
                tabular.read_workbook(long_cell)
            assert caught.value.kind == "too-large"
        else:
            assert len(tabular.read_workbook(long_cell).rows[0][0]) == tabular.MAX_CELL


def _with_sheet_xml(data: bytes, rows_xml: str, dimension: str) -> bytes:
    """The workbook with its first sheet's cells replaced by hand-written XML."""
    with zipfile.ZipFile(io.BytesIO(data)) as zf:
        parts = {n: zf.read(n) for n in zf.namelist()}
    parts["xl/worksheets/sheet1.xml"] = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
        f'<dimension ref="{dimension}"/><sheetData>{rows_xml}</sheetData></worksheet>'
    ).encode()
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zf:
        for name, raw in parts.items():
            zf.writestr(name, raw)
    return out.getvalue()


def test_a_far_cell_is_refused_or_read_without_filling_the_gap():
    # openpyxl's row iterator pads to the sheet's dimension, so one cell in
    # XFD1048576 would make millions of empty values; read through its parser,
    # nothing between is made.
    base = workbook(lambda book: book.active.append(["id"]))
    far_column = _with_sheet_xml(
        base, '<row r="1"><c r="A1" t="inlineStr"><is><t>id</t></is></c></row>'
              '<row r="1048576"><c r="XFD1048576"><v>1</v></c></row>', "A1:XFD1048576")
    start = time.perf_counter()
    with pytest.raises(tabular.TabularError, match="value in column 16,384"):
        tabular.read_workbook(far_column)
    far_row = _with_sheet_xml(
        base, '<row r="1"><c r="A1" t="inlineStr"><is><t>id</t></is></c></row>'
              '<row r="1048576"><c r="A1048576"><v>7</v></c></row>', "A1:XFD1048576")
    table = tabular.read_workbook(far_row)
    assert table.rows == [["7"]] and table.total == 1
    assert time.perf_counter() - start < 5


def test_cell_text():
    assert tabular.cell_text(None) == ""
    assert tabular.cell_text(12) == "12"
    assert tabular.cell_text(12.0) == "12"
    assert tabular.cell_text(-0.5) == "-0.5"
    assert tabular.cell_text(1e-05) == "0.00001"
    assert tabular.cell_text(1e20) == "100000000000000000000"
    assert tabular.cell_text(True) == "true"
    assert tabular.cell_text(datetime.date(2020, 2, 29)) == "2020-02-29"
    assert tabular.cell_text(datetime.datetime(2020, 2, 29)) == "2020-02-29"
    assert tabular.cell_text(datetime.datetime(2020, 2, 29, 8, 5, 3)) == "2020-02-29T08:05:03"
    assert tabular.cell_text(datetime.time(8, 5)) == "08:05:00"
    assert tabular.cell_text(datetime.timedelta(hours=1, minutes=30)) == "PT1H30M0S"
    assert tabular.cell_text("#DIV/0!") == "#DIV/0!"


def test_a_workbook_is_told_from_csv_by_name_or_first_bytes():
    data = workbook(orgs_book)
    assert tabular.is_workbook(data, "orgs.xlsx")
    # A workbook named .csv is still not CSV text.
    assert tabular.is_workbook(data, "orgs.csv")
    assert tabular.is_workbook(data, "")
    assert not tabular.is_workbook(b"id,name\n1,Ann\n", "people.csv")
    assert not tabular.is_workbook(b"id,name\n1,Ann\n", "")
    assert tabular.is_workbook(b"id,name\n", "book.xlsm")
    assert tabular.read_file(b"id;name\n1;Ann\n", "people.csv", {}).separator == ";"


@pytest.mark.perf
def test_read_a_5000_row_sheet_budget():
    # Section 10: a 5,000-row sheet in 1 s.
    def build(book):
        sheet = book.active
        sheet.append(["id", "name", "born", "staff", "city"])
        for i in range(5000):
            sheet.append([f"id{i}", f"Somebody {i}", datetime.date(1990, 1, 1 + i % 28), i, "Somewhere"])

    data = workbook(build)
    times = []
    gc.disable()
    try:
        for _ in range(5):
            start = time.perf_counter()
            tabular.inspect(tabular.read_file(data, "big.xlsx", {}))
            times.append((time.perf_counter() - start) * 1000)
    finally:
        gc.enable()
    median = sorted(times)[2]
    assert median <= limit_ms(1000), f"reading a 5,000-row sheet took {median:.0f} ms (median of 5)"
